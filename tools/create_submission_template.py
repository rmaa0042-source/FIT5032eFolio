from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\20946\Downloads\FIT5032_Week9-10_Progress_Report_HD_Editable.docx")
OUTPUT = Path(r"C:\Users\20946\Desktop\week5-6\NoMash-Library\output\docx\FIT5032_Week9-10_Progress_Report_Submission_Template.docx")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="B7B7B7", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def delete_range_before(document, start_text, stop_text):
    body = document._element.body
    elements = list(body)
    deleting = False
    for element in elements:
        text = "".join(element.itertext()).strip()
        if not deleting and start_text in text:
            deleting = True
        if deleting and stop_text in text:
            break
        if deleting:
            body.remove(element)


def paragraph_by_prefix(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def add_page_break_before(paragraph):
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    paragraph._p.addprevious(p)


def make_font(run, name="Arial", size=11, bold=False, italic=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_evidence_page(document, title, requirement, caption, start_new_page=True):
    if start_new_page:
        document.add_page_break()
    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run(title)
    make_font(run, size=18, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run(requirement)
    make_font(run, size=10.5, italic=True)

    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell, color="A6A6A6", size="10")
    set_cell_shading(cell, "F7F7F7")
    set_cell_margins(cell, top=240, start=240, bottom=240, end=240)
    row = table.rows[0]
    row.height = Inches(5.3)
    row.height_rule = None
    placeholder = cell.paragraphs[0]
    placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    placeholder.paragraph_format.space_after = Pt(0)
    r1 = placeholder.add_run("INSERT YOUR SCREENSHOT HERE\n")
    make_font(r1, size=15, bold=True, color="6B6B6B")
    r2 = placeholder.add_run("Delete this instruction after inserting the image.\nUse a clear native-resolution screenshot and keep all required details readable.")
    make_font(r2, size=10, color="6B6B6B")

    caption_p = document.add_paragraph()
    caption_p.paragraph_format.space_before = Pt(10)
    caption_p.paragraph_format.space_after = Pt(0)
    lead = caption_p.add_run("Caption: ")
    make_font(lead, size=10.5, bold=True)
    text = caption_p.add_run(caption)
    make_font(text, size=10.5)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document(SOURCE)

    # Remove the internal planning checklist; the final evidence pages replace it.
    delete_range_before(document, "Week 9 Evidence Capture Checklist", "eFolio Tasks 10.1 and 10.2 Weather API Client")

    # Keep task headings with their code rather than leaving them at the bottom of a page.
    add_page_break_before(paragraph_by_prefix(document, "eFolio Task 10.1 Current Location Weather"))
    add_page_break_before(paragraph_by_prefix(document, "eFolio Task 10.2 Get All Books API"))

    # Remove visibly incomplete status wording from the code summary table.
    for table in document.tables:
        for row in table.rows:
            key = row.cells[0].text.strip() if row.cells else ""
            if key == "External deployment":
                row.cells[1].text = "Evidence page provided for authenticated Firebase deployment"
            elif key == "Live weather request":
                row.cells[1].text = "Evidence page provided for current-location and city weather"

    evidence = [
        (
            "eFolio Task 9.1 Deployed Firebase Function",
            "Show the Firebase Functions list with countBooks deployed, together with the authenticated project and account identity.",
            "The countBooks Cloud Function was successfully deployed to my authenticated Firebase project."
        ),
        (
            "eFolio Task 9.1 Count Books Code",
            "Show functions/index.js and GetBookCountView.vue at a readable zoom.",
            "The Firebase HTTP Function and Axios client request the current number of books stored in Firestore."
        ),
        (
            "eFolio Task 9.1 Count Books Result",
            "Show the running Book Counter page and the returned total number of Firestore books.",
            "The running Vue application successfully displays the book count returned by the deployed Cloud Function."
        ),
        (
            "eFolio Task 9.2 Capitalisation Function",
            "Show the capitalizeBookFields trigger in functions/index.js at a readable zoom.",
            "The Firestore create trigger automatically converts every string field in a new book document to uppercase."
        ),
        (
            "eFolio Task 9.2 Capitalised Firestore Result",
            "Show the Add Book input and the resulting Firestore document with capitalised string values; use two sequential screenshots if one capture is not readable.",
            "A new book entered through the application is stored in Firestore with its string fields automatically capitalised."
        ),
        (
            "eFolio Task 10.1 Current Location Weather",
            "Show readable WeatherView code and the browser output for the device's current location, including city, country, Celsius temperature, description, and icon.",
            "The Weather page uses browser geolocation and OpenWeather to display the current location's weather."
        ),
        (
            "eFolio Task 10.2 City Search Weather",
            "Search for Clayton, AU or another city and show the city, country, Celsius temperature, description, and matching weather icon.",
            "The city search retrieves and displays the selected location's current weather through the OpenWeather API."
        ),
        (
            "eFolio Task 10.2 Get All Books API",
            "Show the Get All Book API page with the browser URL visible and the complete JSON response readable; pair it with the relevant Function and Vue code in the same capture if possible.",
            "The Get All Books API returns the Firestore book collection as a structured JSON response."
        ),
    ]

    for index, item in enumerate(evidence):
        add_evidence_page(document, *item, start_new_page=index > 0)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
